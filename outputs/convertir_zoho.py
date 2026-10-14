#!/usr/bin/env python3
"""Conversion locale du deuxième onglet d'un planning vers le modèle Zoho.

Dépendance : pip install openpyxl
Aucune connexion réseau. Aucun import dans Zoho.
"""
import argparse
import json
import re
from copy import copy
from datetime import datetime, date
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment


def convert(source, template, output, sheet=None, list_name="Migration Linux 2026-2027", list_id=None, owners=None, status=None, overrides=None):
    if output.resolve() in {source.resolve(), template.resolve()}:
        raise ValueError("Le fichier de sortie doit être distinct des sources.")
    src = openpyxl.load_workbook(source, data_only=True)
    ws = src[sheet] if sheet else src.worksheets[1]
    header_row = next((r for r in range(1, ws.max_row + 1)
                       if ws.cell(r, 1).value == "ID" and ws.cell(r, 2).value == "Action"), None)
    if header_row is None:
        raise ValueError("En-tête ID / Action introuvable.")
    headers = {str(c.value).strip(): c.column for c in ws[header_row] if c.value}
    required = ["ID", "Action", "Début proposé", "Fin proposée", "Pilote proposé / appuis", "Prédécesseurs", "Statut à renseigner", "Critère d’acceptation / prérequis", "Source / écart à l’ébauche"]
    if any(h not in headers for h in required):
        raise ValueError("Colonnes source manquantes : " + ", ".join(h for h in required if h not in headers))
    tasks = []
    for row in range(header_row + 1, ws.max_row + 1):
        identifier = ws.cell(row, headers["ID"]).value
        if not identifier:
            break  # Fin du tableau de tâches, avant le référentiel CDB/PDB.
        if not re.fullmatch(r"[A-Za-z]+\d+", str(identifier)):
            raise ValueError(f"Identifiant inattendu ligne {row} : {identifier}")
        task = {h: ws.cell(row, c).value for h, c in headers.items()}
        task["ID"] = str(identifier)
        tasks.append(task)
    if not tasks:
        raise ValueError("Aucune tâche trouvée.")
    by_id = {t["ID"]: t for t in tasks}
    if len(by_id) != len(tasks):
        raise ValueError("Identifiants de tâches dupliqués.")
    wb = openpyxl.load_workbook(template)
    out = wb.worksheets[0]
    # Un seul onglet d'import : pas d'exemples ni de tableaux annexes.
    for other in list(wb.worksheets)[1:]:
        wb.remove(other)
    target_headers = [c.value for c in out[1]]
    styles = [copy(c._style) for c in out[2]]
    for name in ["TASK NAME", "TASK ID", "PREDECESSORS", "TASK DESCRIPTION", "TASK LIST NAME", "START DATE", "DUE DATE"]:
        if name not in target_headers:
            raise ValueError(f"Colonne absente du modèle : {name}")
    out.delete_rows(2, out.max_row)
    # Champ personnalisé RACI, à associer dans l'assistant d'import Zoho.
    if "Responsible" not in target_headers:
        target_headers.append("Responsible")
        out.cell(1, len(target_headers), "Responsible")._style = copy(out.cell(1, 1)._style)
        styles.append(copy(styles[0]))
    warnings = []
    graph = {}
    for task in tasks:
        identifier = task["ID"]
        start, end = task["Début proposé"], task["Fin proposée"]
        if not isinstance(start, (datetime, date)) or not isinstance(end, (datetime, date)) or end < start:
            raise ValueError(f"Dates invalides pour {identifier}")
        name = str(task["Action"] or "").strip()
        if not name or len(name) > 255:
            raise ValueError(f"Nom vide ou trop long pour {identifier}")
        raw = str(task["Prédécesseurs"] or "").strip()
        deps = [] if raw in {"", "—", "-"} else [p.strip() for p in raw.split(",")]
        relations = [(p, "FS") for p in deps]
        # L'exception est décrite explicitement dans le détail de la tâche.
        if "L17" in by_id and identifier == "L18" and "début de L17" in str(task["Source / écart à l’ébauche"]):
            relations.append(("L17", "SS"))
        if overrides and identifier in overrides:
            relations = []
            for entry in overrides[identifier]:
                match = re.fullmatch(r"([A-Za-z]+\d+)\s+(FS|SS|FF|SF)", entry)
                if not match:
                    raise ValueError(f"Dépendance invalide : {entry}")
                relations.append(match.groups())
        graph[identifier] = [p for p, _ in relations]
        for pred, relation in relations:
            if pred not in by_id or pred == identifier:
                raise ValueError(f"Prédécesseur invalide pour {identifier} : {pred}")
            prev = by_id[pred]
            if relation == "FS" and start <= prev["Fin proposée"]:
                warnings.append(f"{identifier} / {pred} FS : début {start:%d/%m/%Y}, fin du prédécesseur {prev['Fin proposée']:%d/%m/%Y}. Dates conservées, chevauchement ou même jour à vérifier.")
            if relation == "SS" and start < prev["Début proposé"]:
                warnings.append(f"{identifier} / {pred} SS : début antérieur au prédécesseur.")
        description = "\n\n".join([
            f"ID source : {identifier}",
            f"Pilote proposé / appuis : {task['Pilote proposé / appuis'] or ''}",
            f"Statut source : {task['Statut à renseigner'] or ''}",
            f"Critère d’acceptation / prérequis : {task['Critère d’acceptation / prérequis'] or ''}",
            f"Source / écart à l’ébauche : {task['Source / écart à l’ébauche'] or ''}",
            "Dates proposées à confirmer. La durée calendaire n’est pas une charge de travail.",
        ])
        values = {"TASK NAME": name, "TASK ID": identifier,
                  "PREDECESSORS": ", ".join(f"{p} {rel}" for p, rel in relations) or None,
                  "TASK DESCRIPTION": description, "TASK LIST NAME": str(list_id) if list_id else list_name,
                  "START DATE": start, "DUE DATE": end,
                  "TASK OWNER": (owners or {}).get(identifier), "CUSTOM STATUS": status,
                  "Responsible": task["Pilote proposé / appuis"]}
        out.append([values.get(str(h).strip()) for h in target_headers])
        for c, style in zip(out[out.max_row], styles):
            c._style = copy(style)
            if c.column in [target_headers.index("START DATE") + 1, target_headers.index("DUE DATE") + 1]:
                c.number_format = "dd/mm/yyyy"
    visited, active = set(), set()
    def visit(identifier):
        if identifier in active:
            raise ValueError(f"Cycle de dépendances : {identifier}")
        if identifier in visited:
            return
        active.add(identifier)
        for pred in graph[identifier]:
            visit(pred)
        active.remove(identifier)
        visited.add(identifier)
    for identifier in graph:
        visit(identifier)
    out.title = "Import Zoho"
    out.freeze_panes = "C2"
    out.auto_filter.ref = out.dimensions
    out.column_dimensions["A"].width = 65
    out.column_dimensions["B"].width = 12
    out.column_dimensions["C"].width = 45
    out.column_dimensions["G"].width = 85
    out.column_dimensions["I"].width = 35
    out.column_dimensions["M"].width = 15
    out.column_dimensions["N"].width = 15
    responsible_column = target_headers.index("Responsible") + 1
    out.column_dimensions[openpyxl.utils.get_column_letter(responsible_column)].width = 50
    for row in out.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
        row[responsible_column - 1].alignment = Alignment(wrap_text=True, vertical="top")
        out.row_dimensions[row[0].row].height = 180
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    saved = openpyxl.load_workbook(output, data_only=True)
    assert saved.active.max_row == len(tasks) + 1
    assert [c.value for c in saved.active[1]] == target_headers
    report = output.with_suffix(".rapport.txt")
    report.write_text(f"{len(tasks)} tâches converties depuis {ws.title}.\n"
                      "Pilotes copiés intégralement dans Responsible (RACI) et conservés dans la description. Propriétaires non attribués sauf mapping explicite.\n"
                      "Avancement et charge laissés vides. Dates conservées.\n"
                      "L18 → L17 SS ajouté lorsque l'exception est présente dans la source.\n"
                      + ("\nPoints à vérifier :\n" + "\n".join(warnings) if warnings else "\nPas de conflit de dates détecté sur les liaisons FS/SS.\n"), encoding="utf-8")
    print(f"{len(tasks)} tâches → {output}\nRapport → {report}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--sheet", help="Nom de l'onglet source (par défaut : deuxième onglet)")
    parser.add_argument("--list-name", default="Migration Linux 2026-2027")
    parser.add_argument("--list-id", help="ID existant : écrit dans TASK LIST NAME selon la documentation Zoho")
    parser.add_argument("--owners", type=Path, help='JSON par ID de tâche : {"L01": "utilisateur Zoho"}')
    parser.add_argument("--status", help="Statut Zoho confirmé, sinon champ vide")
    parser.add_argument("--dependencies", type=Path, help='JSON de remplacements : {"L18": ["L17 SS", "L05 FS"]}')
    args = parser.parse_args()
    def load(path):
        return json.loads(path.read_text(encoding="utf-8")) if path else None
    try:
        convert(args.source, args.template, args.output, args.sheet, args.list_name,
                args.list_id, load(args.owners), args.status, load(args.dependencies))
    except (ValueError, KeyError, IndexError) as exc:
        parser.exit(1, f"Erreur : {exc}\n")


if __name__ == "__main__":
    main()
